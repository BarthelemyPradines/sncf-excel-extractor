"""Data classes, shared constants, and utility helpers."""

from __future__ import annotations

import re
from dataclasses import dataclass, field


@dataclass
class SubTable:
    top: int
    left: int
    bottom: int
    right: int
    header_rows: int
    cell_count: int

    def excel_range(self) -> str:
        from openpyxl.utils import get_column_letter
        return (f"{get_column_letter(self.left)}{self.top}"
                f":{get_column_letter(self.right)}{self.bottom}")


@dataclass
class SheetReport:
    name: str
    classification: str
    n_rows: int
    n_cols: int
    n_filled: int
    n_formulas: int
    n_cross_sheet_refs: int
    n_merged: int
    n_charts: int
    n_images: int
    has_listobjects: bool
    sub_tables: list[SubTable] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class WorkbookReport:
    path: str
    sheets: list[SheetReport]
    extracted_media: list[str]
    extracted_charts: list[str]
    values_by_sheet: dict[str, list[list]] = field(default_factory=dict, repr=False)


# -- shared constants --

CROSS_SHEET_RE = re.compile(r"(?:'[^']+'|[A-Za-z_][\w\.]*)!")
SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r":    "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


# -- shared helpers --

def safe_name(name: str) -> str:
    """Sanitize a sheet name for use as a filename."""
    return SAFE_NAME_RE.sub("_", name).strip("_") or "sheet"


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def normalize_target(target: str) -> str:
    """Resolve relative xlsx zip paths like '../drawings/drawing1.xml'
    or absolute ones like '/xl/drawings/drawing1.xml' to 'xl/...'."""
    target = target.lstrip("/")
    if target.startswith("xl/"):
        return target
    if target.startswith("../"):
        return "xl/" + target[3:]
    return "xl/" + target
