"""Top-level workbook analysis orchestrator."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .models import SubTable, SheetReport, WorkbookReport
from .scanner import scan_sheet
from .detection import split_by_blank_lines, guess_header_rows, classify
from .media import extract_media, zip_metadata

try:
    import formulas as _formulas_mod  # type: ignore
    HAVE_FORMULAS = True
except ImportError:
    HAVE_FORMULAS = False

CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$")


def _fill_missing_with_formulas(xlsx_path: Path,
                                values_by_sheet: dict[str, list[list]]) -> None:
    """Use the `formulas` library to compute missing values in-place."""
    has_holes = False
    for values in values_by_sheet.values():
        for row in values:
            if any(v is None for v in row):
                has_holes = True
                break
        if has_holes:
            break
    if not has_holes:
        return

    try:
        xl = _formulas_mod.ExcelModel().loads(str(xlsx_path)).finish()
        result = xl.calculate()
    except Exception as e:
        print(f"  warning: formula evaluation failed: {e}")
        return

    fname = xlsx_path.name
    for key, ranges_obj in result.items():
        key_str = str(key)
        if fname not in key_str:
            continue
        try:
            sheet_part, cell_ref = key_str.rsplit("!", 1)
            sheet_name = sheet_part.split("]", 1)[1].strip("'")
        except (ValueError, IndexError):
            continue

        m = CELL_REF_RE.match(cell_ref)
        if not m:
            continue
        col = column_index_from_string(m.group(1)) - 1
        row = int(m.group(2)) - 1

        values = values_by_sheet.get(sheet_name)
        if values is None:
            for name in values_by_sheet:
                if name.upper() == sheet_name.upper():
                    values = values_by_sheet[name]
                    break
        if values is None or row >= len(values):
            continue
        if col >= len(values[row]):
            continue

        if values[row][col] is not None:
            continue

        try:
            val = ranges_obj.value
            if hasattr(val, '__iter__') and not isinstance(val, str):
                for item in val:
                    if hasattr(item, '__iter__') and not isinstance(item, str):
                        for v in item:
                            val = v
                            break
                        break
                    else:
                        val = item
                        break
            values[row][col] = val
        except Exception:
            continue


def analyze_workbook(xlsx_path: str | Path,
                     media_dir: str | Path | None = None,
                     formula_results: bool = True) -> WorkbookReport:
    xlsx_path = Path(xlsx_path)
    meta = zip_metadata(xlsx_path)

    wb = load_workbook(xlsx_path, read_only=True, data_only=formula_results)

    sheet_reports: list[SheetReport] = []
    values_by_sheet: dict[str, list[list]] = {}

    for name in wb.sheetnames:
        sheet_meta = meta.get(name, dict(n_merged=0, n_charts=0, n_images=0,
                                         has_listobjects=False))
        sig, mask, values = scan_sheet(wb[name])

        values_by_sheet[name] = values

        boxes = split_by_blank_lines(mask)
        single_region = len(boxes) <= 1

        sub_tables: list[SubTable] = []
        for r0, c0, r1, c1 in boxes:
            top, left, bottom, right = r0 + 1, c0 + 1, r1 + 1, c1 + 1
            cells = (bottom - top + 1) * (right - left + 1)
            if cells < 4:
                continue
            h = 1 if single_region else guess_header_rows(values, r0, c0, c1)
            sub_tables.append(SubTable(top, left, bottom, right, h, cells))

        cls, notes = classify(sig, sheet_meta, len(sub_tables))
        sheet_reports.append(SheetReport(
            name=name, classification=cls,
            n_rows=sig["n_rows"], n_cols=sig["n_cols"],
            n_filled=sig["n_filled"],
            n_formulas=sheet_meta["n_formulas"],
            n_cross_sheet_refs=sheet_meta["n_cross_sheet_refs"],
            n_merged=sheet_meta["n_merged"],
            n_charts=sheet_meta["n_charts"],
            n_images=sheet_meta["n_images"],
            has_listobjects=sheet_meta["has_listobjects"],
            sub_tables=sub_tables, notes=notes,
        ))

    wb.close()

    if formula_results and HAVE_FORMULAS:
        _fill_missing_with_formulas(xlsx_path, values_by_sheet)

    if media_dir is None:
        media_dir = xlsx_path.with_suffix("").as_posix() + "_media"
    images, charts = extract_media(xlsx_path, media_dir)

    return WorkbookReport(
        path=str(xlsx_path),
        sheets=sheet_reports,
        extracted_media=images,
        extracted_charts=charts,
        values_by_sheet=values_by_sheet,
    )
